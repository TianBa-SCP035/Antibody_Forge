import unittest

from openpyxl import load_workbook
from sqlalchemy import BigInteger, create_engine
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import sessionmaker

from db.session import Base
from models.molecular_cell import MolecularLibraryOrder, MolecularPrimerIndexCatalog
from modules.molecular_cell.primer_catalog import service


@compiles(BigInteger, "sqlite")
def _compile_big_integer_for_sqlite(_type, _compiler, **_kwargs):
    return "INTEGER"


class PrimerCatalogServiceTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite+pysqlite:///:memory:")
        Base.metadata.create_all(self.engine, tables=[MolecularPrimerIndexCatalog.__table__])
        self.Session = sessionmaker(bind=self.engine, autoflush=False, expire_on_commit=False)
        self.db = self.Session()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def test_save_normalizes_and_hides_inactive_from_options(self):
        saved = service.save(
            self.db,
            {
                "name": " UDP-1-F ",
                "family": "UDP",
                "direction": "f",
                "version": "V1",
                "short_sequence": "ac gtn",
                "homology_arm_1": "aa",
                "homology_arm_2": "tt",
                "note": "新引物",
            },
        )
        self.assertEqual(saved["name"], "UDP-1-F")
        self.assertEqual(saved["direction"], "F")
        self.assertEqual(saved["short_sequence"], "ACGTN")
        self.assertTrue(saved["active"])
        self.assertEqual(service.list_options(self.db, query="UDP-1")["items"][0]["name"], "UDP-1-F")

        hidden = service.save(
            self.db,
            {
                "id": saved["id"],
                "active": False,
                "name": saved["name"],
                "family": saved["family"],
                "direction": saved["direction"],
                "short_sequence": saved["short_sequence"],
            },
        )
        self.assertFalse(hidden["active"])
        self.assertEqual(hidden["version"], "V1")
        self.assertEqual(hidden["homology_arm_1"], "AA")
        self.assertEqual(hidden["homology_arm_2"], "TT")
        self.assertEqual(service.list_options(self.db, query="UDP-1")["items"], [])

    def test_list_filters_and_families(self):
        service.save(
            self.db,
            {"name": "UDP-1-F", "family": "UDP", "direction": "F", "short_sequence": "AAAA", "active": False},
        )
        service.save(
            self.db,
            {"name": "UDP-2-R", "family": "UDP", "direction": "R", "short_sequence": "TTTT", "note": "第二条"},
        )
        service.save(
            self.db,
            {"name": "IDX-7", "family": "达普Barcode", "direction": "F", "short_sequence": "GGGG"},
        )

        listed = service.list_catalog(self.db, {"name": "UDP", "page": 1, "limit": 1})
        self.assertEqual(listed["total"], 2)
        self.assertEqual(len(listed["items"]), 1)
        self.assertEqual([item["name"] for item in listed["stats"]["families"]], ["达普Barcode", "UDP"])
        self.assertEqual(service.check_sequence("AGGTGCGT"), "ACGCACCT")

        self.assertEqual(service.list_catalog(self.db, {"direction": "R"})["total"], 1)
        self.assertEqual(service.list_catalog(self.db, {"active": False})["total"], 1)
        stats = service.list_catalog(self.db, {"family": "达普Barcode"})["stats"]
        by_name = {item["name"]: item["count"] for item in stats["families"]}
        self.assertEqual(stats["total"], 3)
        self.assertEqual(by_name["达普Barcode"], 1)
        self.assertEqual(by_name["UDP"], 2)
        narrowed = service.list_catalog(self.db, {"direction": "R"})["stats"]
        narrowed_names = {item["name"]: item["count"] for item in narrowed["families"]}
        self.assertEqual(narrowed["total"], 1)
        self.assertEqual(narrowed_names["UDP"], 1)

        custom = service.save(
            self.db,
            {"name": "CUSTOM-1", "family": "自建库", "direction": "F", "short_sequence": "ACGT"},
        )
        self.assertEqual(custom["family"], "自建库")
        custom_names = [item["name"] for item in service.list_catalog(self.db)["stats"]["families"]]
        self.assertEqual(custom_names[:2], ["达普Barcode", "UDP"])
        self.assertIn("自建库", custom_names)

    def test_save_rejects_duplicates_bad_sequences_and_bad_flags(self):
        saved = service.save(
            self.db,
            {"name": "UDP-1-F", "family": "UDP", "direction": "F", "short_sequence": "ACGT"},
        )
        with self.assertRaisesRegex(ValueError, service.NAME_TAKEN):
            service.save(
                self.db,
                {"name": "UDP-1-F", "family": "UDP", "direction": "F", "short_sequence": "AAAA"},
            )
        with self.assertRaisesRegex(ValueError, "A、C、G、T、N"):
            service.save(
                self.db,
                {"name": "坏序列", "family": "UDP", "direction": "F", "short_sequence": "ACGX"},
            )
        with self.assertRaisesRegex(ValueError, "正向或反向"):
            service.save(
                self.db,
                {"name": "坏方向", "family": "UDP", "direction": "X", "short_sequence": "ACGT"},
            )
        with self.assertRaisesRegex(ValueError, "启用状态"):
            service.save(
                self.db,
                {
                    "id": saved["id"],
                    "name": saved["name"],
                    "family": saved["family"],
                    "direction": saved["direction"],
                    "short_sequence": saved["short_sequence"],
                    "active": "false",
                },
            )
        with self.assertRaisesRegex(ValueError, service.MISSING_PRIMER):
            service.save(
                self.db,
                {"id": 999, "name": "X", "family": "UDP", "direction": "F", "short_sequence": "ACGT"},
            )

    def test_check_keyword_and_batch_edit(self):
        Base.metadata.create_all(self.engine, tables=[MolecularLibraryOrder.__table__])
        first = service.save(
            self.db,
            {"name": "BATCH-F", "family": "UDP", "direction": "F", "short_sequence": "AAAA"},
        )
        second = service.save(
            self.db,
            {"name": "BATCH-R", "family": "UDP", "direction": "R", "short_sequence": "CCCC"},
        )
        self.db.add(
            MolecularLibraryOrder(
                library_order_id="WO-BATCH",
                build_type="phage_ngs",
                status="待处理",
                priority="普通",
                i7_catalog_id=second["id"],
            )
        )
        self.db.commit()

        matched = service.list_catalog(self.db, {"check_keyword": "tt"})
        self.assertEqual([item["name"] for item in matched["items"]], ["BATCH-F"])
        self.assertEqual(service.list_ids(self.db, {"check_keyword": "TTTT"}), [first["id"]])
        self.assertEqual(service.list_catalog(self.db, {"check_keyword": "AX"})["total"], 0)

        self.assertEqual(
            service.batch_update(self.db, [first["id"], second["id"]], family="自建系列", active=False)["total"],
            2,
        )
        refreshed = self.db.get(MolecularPrimerIndexCatalog, first["id"])
        self.assertEqual(refreshed.family, "自建系列")
        self.assertFalse(refreshed.active)

        result = service.batch_delete(self.db, [first["id"], second["id"]])
        self.assertEqual(result["deleted"], 1)
        self.assertEqual(result["skipped"], ["BATCH-R"])
        self.assertIsNone(self.db.get(MolecularPrimerIndexCatalog, first["id"]))
        self.assertIsNotNone(self.db.get(MolecularPrimerIndexCatalog, second["id"]))

    def test_delete_unused_primer_and_reject_referenced_one(self):
        Base.metadata.create_all(self.engine, tables=[MolecularLibraryOrder.__table__])
        unused = service.save(
            self.db,
            {"name": "UNUSED-F", "family": "UDP", "direction": "F", "short_sequence": "AAAA"},
        )
        used = service.save(
            self.db,
            {"name": "USED-R", "family": "UDP", "direction": "R", "short_sequence": "TTTT"},
        )
        self.db.add(
            MolecularLibraryOrder(
                library_order_id="WO-1",
                build_type="phage_ngs",
                status="待处理",
                priority="普通",
                h_forward_primer_id=used["id"],
            )
        )
        self.db.commit()

        self.assertEqual(service.delete_primer(self.db, unused["id"])["id"], unused["id"])
        self.assertIsNone(self.db.get(MolecularPrimerIndexCatalog, unused["id"]))
        with self.assertRaisesRegex(ValueError, service.PRIMER_IN_USE):
            service.delete_primer(self.db, used["id"])
        with self.assertRaisesRegex(ValueError, service.MISSING_PRIMER):
            service.delete_primer(self.db, unused["id"])

    def test_batch_upserts_by_name_and_exports(self):
        existing = service.save(
            self.db,
            {
                "name": "UDP-1-F",
                "family": "UDP",
                "direction": "F",
                "short_sequence": "AAAA",
                "version": "V1",
                "homology_arm_1": "AA",
                "homology_arm_2": "TT",
                "active": False,
            },
        )
        result = service.batch_save(
            self.db,
            [
                {"name": "UDP-1-F", "family": "UDP", "direction": "F", "short_sequence": "CCCC"},
                {
                    "name": "UDP-2-R",
                    "family": "UDP",
                    "direction": "R",
                    "short_sequence": "GGGG",
                    "active": True,
                },
            ],
        )
        self.assertEqual(result, {"total": 2, "created": 1, "updated": 1})
        updated = self.db.get(MolecularPrimerIndexCatalog, existing["id"])
        self.assertEqual(updated.short_sequence, "CCCC")
        self.assertFalse(updated.active)
        self.assertEqual(updated.version, "V1")
        self.assertEqual(updated.homology_arm_1, "AA")
        self.assertEqual(updated.homology_arm_2, "TT")

        output, _filename = service.export_workbook(self.db, {"direction": "R"})
        workbook = load_workbook(output, read_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()
        self.assertEqual(rows[0], service.EXPORT_HEADERS)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "UDP-2-R")
        self.assertEqual(rows[1][2], "反向")
        self.assertEqual(rows[1][4], "CCCC")

    def test_batch_rolls_back_as_one_transaction(self):
        existing = service.save(
            self.db,
            {"name": "UDP-1-F", "family": "UDP", "direction": "F", "short_sequence": "AAAA"},
        )
        with self.assertRaisesRegex(ValueError, "第 2 行"):
            service.batch_save(
                self.db,
                [
                    {"name": "OK", "family": "UDP", "direction": "F", "short_sequence": "CCCC"},
                    {"name": "BAD", "family": "UDP", "direction": "F", "short_sequence": "CCCX"},
                ],
            )
        unchanged = self.db.get(MolecularPrimerIndexCatalog, existing["id"])
        self.assertEqual(unchanged.short_sequence, "AAAA")
        with self.assertRaisesRegex(ValueError, "名称重复"):
            service.batch_save(
                self.db,
                [
                    {"name": "DUP", "family": "UDP", "direction": "F", "short_sequence": "CCCC"},
                    {"name": "DUP", "family": "UDP", "direction": "F", "short_sequence": "GGGG"},
                ],
            )

if __name__ == "__main__":
    unittest.main()
