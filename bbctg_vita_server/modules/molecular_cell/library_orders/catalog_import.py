"""Import the experimenter's primer/Index workbook into the catalog table."""

from argparse import ArgumentParser
from collections import Counter
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from db.session import SessionLocal
from models.molecular_cell import MolecularPrimerIndexCatalog


LEGACY_FAMILY = "legacy_8nt"
UDP_FAMILY = "udp_10nt"
DNA_PATTERN = re.compile(r"^[ACGTN]+$", re.IGNORECASE)
DIRECTION_PATTERN = re.compile(r"-(F|R)$", re.IGNORECASE)
VERSION_PATTERN = re.compile(r"(V\d+)", re.IGNORECASE)


def _text(value) -> str:
    return str(value or "").strip()


def _sequence(value, *, row: int, column: str) -> str:
    sequence = _text(value).upper()
    if not sequence or not DNA_PATTERN.fullmatch(sequence):
        raise ValueError(f"{column}{row} 不是有效DNA短序列")
    return sequence


def _template(sheet, label_row: int, value_row: int) -> tuple[str | None, str | None]:
    if _text(sheet.cell(label_row, 11).value).endswith("同源臂1"):
        arm_1 = _sequence(sheet.cell(value_row, 11).value, row=value_row, column="K")
        arm_2 = _sequence(sheet.cell(value_row, 13).value, row=value_row, column="M")
        return arm_1, arm_2
    return None, None


def parse_catalog_workbook(path: str | Path) -> list[dict]:
    workbook_path = Path(path)
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        forward_arms = _template(sheet, 1, 2)
        reverse_arms = _template(sheet, 4, 5)
        items: list[dict] = []

        for row in range(1, sheet.max_row + 1):
            legacy_name = _text(sheet.cell(row, 1).value)
            legacy_sequence = _text(sheet.cell(row, 2).value)
            if legacy_name or legacy_sequence:
                items.append(
                    {
                        "name": legacy_name,
                        "family": LEGACY_FAMILY,
                        "direction": None,
                        "version": None,
                        "short_sequence": _sequence(legacy_sequence, row=row, column="B"),
                        "homology_arm_1": None,
                        "homology_arm_2": None,
                        "source_file": workbook_path.name,
                        "source_row": row,
                    }
                )

            udp_name = _text(sheet.cell(row, 6).value)
            udp_sequence = _text(sheet.cell(row, 7).value)
            if udp_name or udp_sequence:
                direction_match = DIRECTION_PATTERN.search(udp_name)
                direction = direction_match.group(1).upper() if direction_match else None
                version_match = VERSION_PATTERN.search(udp_name)
                version = version_match.group(1).upper() if version_match else None
                arm_1, arm_2 = forward_arms if direction == "F" else reverse_arms if direction == "R" else (None, None)
                short_sequence = _sequence(udp_sequence, row=row, column="G")
                items.append(
                    {
                        "name": udp_name,
                        "family": UDP_FAMILY,
                        "direction": direction,
                        "version": version,
                        "short_sequence": short_sequence,
                        "homology_arm_1": arm_1,
                        "homology_arm_2": arm_2,
                        "source_file": workbook_path.name,
                        "source_row": row,
                    }
                )

        names = [item["name"] for item in items]
        if any(not name for name in names):
            raise ValueError("字典记录名称不能为空")
        duplicates = sorted(name for name, count in Counter(names).items() if count > 1)
        if duplicates:
            raise ValueError(f"字典名称重复：{', '.join(duplicates[:5])}")
        return items
    finally:
        workbook.close()


def import_catalog(db: Session, path: str | Path) -> dict[str, int]:
    items = parse_catalog_workbook(path)
    names = [item["name"] for item in items]
    existing = {
        row.name: row
        for row in db.scalars(
            select(MolecularPrimerIndexCatalog).where(MolecularPrimerIndexCatalog.name.in_(names))
        ).all()
    }
    created = 0
    updated = 0
    for item in items:
        record = existing.get(item["name"])
        if record is None:
            record = MolecularPrimerIndexCatalog(name=item["name"])
            db.add(record)
            created += 1
        else:
            updated += 1
        for key, value in item.items():
            setattr(record, key, value)
        record.active = True
        record.note = "按来源工作簿原值导入；未推断F/R与i7/i5映射"
    db.commit()
    return {"total": len(items), "created": created, "updated": updated}


def main() -> None:
    parser = ArgumentParser(description="导入文库构建引物/Index字典")
    parser.add_argument("workbook", help="八口命名.xlsx 路径")
    args = parser.parse_args()
    with SessionLocal() as db:
        result = import_catalog(db, args.workbook)
    print(
        f"导入完成：共 {result['total']} 条，新增 {result['created']} 条，更新 {result['updated']} 条"
    )


if __name__ == "__main__":
    main()
